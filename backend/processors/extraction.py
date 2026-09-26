from __future__ import annotations

import csv
import os
import re
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePath

from utils.config import get_settings

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".json", ".html", ".htm", ".xml", ".log", ".yaml", ".yml"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv"} | TEXT_EXTENSIONS



_SIGNATURES: tuple[tuple[bytes, str], ...] = (
    (b"%PDF-", "pdf"),
    (b"\x89PNG\r\n\x1a\n", "image"),
    (b"\xff\xd8\xff", "image"),  
    (b"RIFF", "image"),  
    (b"BM", "image"),  
    (b"II*\x00", "image"),  
    (b"MM\x00*", "image"),  
    (b"GIF87a", "image"),  
    (b"GIF89a", "image"),  
    (b"PK\x03\x04", "zip"),  
)




_MIN_MEANINGFUL_CHARS = 10
_MIN_ALNUM_RATIO = 0.30


class ExtractionError(Exception):

    def __init__(self, message: str, code: str = "UNREADABLE_DOCUMENT"):
        super().__init__(message)
        self.code = code


@dataclass
class PageResult:
    page_number: int
    text: str
    extraction_method: str  
    ocr: bool = False
    warning: str | None = None


@dataclass
class ExtractionResult:
    filename: str
    file_type: str
    success: bool
    partial: bool = False
    analyzable: bool = True
    status: str = "SUCCESS"  
    text: str = ""
    pages: list[PageResult] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    error: str | None = None
    error_code: str | None = None
    extraction_method: str = "native"
    confidence: float = 1.0
    truncated: bool = False


def safe_filename(filename: str | None) -> str:
    return PurePath(filename or "upload").name or "upload"


def file_extension(filename: str | None) -> str:
    return PurePath(safe_filename(filename)).suffix.lower()


def sniff_content_kind(data: bytes) -> str | None:
    head = data[:16]
    if head.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return "image"
    for signature, kind in _SIGNATURES:
        if head.startswith(signature):
            return kind
    return None


def _is_safe_text(data: bytes) -> bool:
    if not data or len(data) > 10 * 1024 * 1024:
        return False
    if b"\x00" in data[:4096]:
        return False
    try:
        sample = data[:8192].decode("utf-8")
    except UnicodeDecodeError:
        try:
            sample = data[:8192].decode("latin-1")
        except Exception:
            return False
    printable = sum(1 for c in sample if c.isprintable() or c in "\r\n\t")
    return (printable / max(1, len(sample))) > 0.80


def detect_file_type(filename: str | None, data: bytes) -> str:
    ext = file_extension(filename)
    sniffed = sniff_content_kind(data)

    if sniffed == "pdf":
        return "pdf"
    if sniffed == "image":
        return "image"
    if sniffed == "zip":
        
        if ext in {".docx", ".pptx", ".xlsx"}:
            return ext[1:]
        try:
            with zipfile.ZipFile(BytesIO(data)) as zf:
                names = set(zf.namelist())
                if any(n.startswith("word/") for n in names):
                    return "docx"
                if any(n.startswith("ppt/") for n in names):
                    return "pptx"
                if any(n.startswith("xl/") for n in names):
                    return "xlsx"
        except Exception:
            pass
        raise ExtractionError("This file type is not supported.", "UNSUPPORTED_FILE_TYPE")

    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in {".docx", ".pptx", ".xlsx", ".csv"}:
        return ext[1:]
    if ext in TEXT_EXTENSIONS:
        return "txt"
    if _is_safe_text(data):
        return "txt"
    raise ExtractionError("This file type is not supported.", "UNSUPPORTED_FILE_TYPE")


def _alnum_ratio(text: str) -> float:
    stripped = text.strip()
    if not stripped:
        return 0.0
    alnum = sum(character.isalnum() or character.isspace() for character in stripped)
    return alnum / len(stripped)


def text_quality_score(text: str) -> float:
    if not text or len(text.strip()) < _MIN_MEANINGFUL_CHARS:
        return 0.0
    ratio = _alnum_ratio(text)
    replacement_junk = text.count("\ufffd") / max(len(text), 1)
    score = ratio - replacement_junk
    return max(0.0, min(1.0, score))


def is_meaningful(text: str) -> bool:
    return text_quality_score(text) >= _MIN_ALNUM_RATIO


def _limit_text(text: str) -> tuple[str, bool]:
    limit = get_settings().max_extracted_text_chars
    if len(text) <= limit:
        return text, False
    return text[:limit] + "\n\n[Text truncated for safe analysis]", True






def _load_image(data: bytes):
    from PIL import Image, ImageOps, UnidentifiedImageError

    try:
        with Image.open(BytesIO(data)) as source:
            source.verify()
        with Image.open(BytesIO(data)) as source:
            image = ImageOps.exif_transpose(source)
            if image.mode in ("RGBA", "LA"):
                background = Image.new("RGB", image.size, "white")
                background.paste(image, mask=image.getchannel("A") if "A" in image.getbands() else None)
                image = background
            else:
                image = image.convert("RGB")
            if min(image.size) and min(image.size) < 800:
                ratio = 800 / max(1, min(image.size))
                image = image.resize((max(1, int(image.width * ratio)), max(1, int(image.height * ratio))))
            return image
    except (UnidentifiedImageError, OSError, ValueError) as exc:
        raise ExtractionError("This image could not be read.", "CORRUPTED_IMAGE") from exc


def local_ocr_available() -> bool:
    try:
        import pytesseract

        command = get_settings().tesseract_cmd
        if command:
            pytesseract.pytesseract.tesseract_cmd = command
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_image_local(image) -> tuple[str | None, bool]:
    try:
        import pytesseract

        command = get_settings().tesseract_cmd
        if command:
            pytesseract.pytesseract.tesseract_cmd = command
        text = pytesseract.image_to_string(image, timeout=get_settings().ocr_timeout_seconds).strip()
        return text, True
    except ImportError:
        return None, False
    except RuntimeError:
        
        return None, False
    except Exception:
        return None, False


def _image_to_png_bytes(image) -> bytes:
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


async def _ocr_image_with_fallback(image, warnings: list[str]) -> tuple[str, str, str]:
    from services.openai_provider import is_configured, vision_ocr

    local_text, local_ok = _ocr_image_local(image)
    if local_ok and local_text and is_meaningful(local_text):
        return local_text, "ocr_local", "SUCCESS"

    if not local_ok:
        warnings.append("Local OCR was unavailable; attempting AI fallback where configured.")
    else:
        warnings.append("Local OCR produced low-confidence text; attempting AI fallback where configured.")

    if is_configured():
        ai_text = await vision_ocr(_image_to_png_bytes(image), "image/png")
        if ai_text and is_meaningful(ai_text):
            return ai_text, "ocr_openai", "SUCCESS"
        warnings.append("AI OCR fallback did not return usable text.")
        return "", "ocr_openai", "NO_ANALYZABLE_CONTENT"

    if local_ok:
        
        return (local_text or ""), "ocr_local", "NO_ANALYZABLE_CONTENT"

    return "", "failed", "OCR_UNAVAILABLE"


async def _extract_image_async(data: bytes) -> tuple[str, str, list[str], float, str]:
    image = _load_image(data)
    warnings: list[str] = []
    text, method, status = await _ocr_image_with_fallback(image, warnings)
    if status == "OCR_UNAVAILABLE":
        raise ExtractionError("OCR could not extract any readable text because text recognition is unavailable.", "OCR_UNAVAILABLE")
    if status == "NO_ANALYZABLE_CONTENT" or not text.strip():
        return "", method, warnings, 0.0, "NO_ANALYZABLE_CONTENT"
    confidence = 1.0 if method == "ocr_local" and is_meaningful(text) else (0.75 if method == "ocr_openai" else 0.5)
    return text, method, warnings, confidence, "SUCCESS"






def _render_pdf_page_to_image(data: bytes, page_index: int):
    try:
        import fitz  
    except ImportError:
        return None
    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            from PIL import Image

            return Image.open(BytesIO(pixmap.tobytes("png")))
    except Exception:
        return None


async def _extract_pdf_async(data: bytes) -> tuple[str, str, list[str], float, dict, list[PageResult]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF extraction is unavailable on this server.", "SERVICE_UNAVAILABLE") from exc

    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This document appears to be corrupted.", "CORRUPTED_DOCUMENT") from exc

    if reader.is_encrypted:
        try:
            
            
            if reader.decrypt("") == 0:
                raise ExtractionError("This PDF is password-protected and cannot be processed.", "ENCRYPTED_DOCUMENT")
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("This PDF is password-protected and cannot be processed.", "ENCRYPTED_DOCUMENT") from exc

    settings = get_settings()
    page_count = len(reader.pages)
    if page_count > settings.max_pdf_pages:
        raise ExtractionError(f"This PDF exceeds the {settings.max_pdf_pages}-page limit.", "DOCUMENT_TOO_LARGE")
    if page_count == 0:
        raise ExtractionError("This PDF has no pages.", "EMPTY_DOCUMENT")

    warnings: list[str] = []
    pages_native, pages_ocr, pages_failed = 0, 0, 0
    results: list[PageResult] = []
    can_rasterize = None  
    ocr_unavailable_pages = 0
    no_readable_pages = 0

    for index in range(page_count):
        try:
            page = reader.pages[index]
            native_text = (page.extract_text() or "").strip()
        except Exception:
            native_text = ""

        if is_meaningful(native_text):
            results.append(PageResult(index + 1, native_text, "native"))
            pages_native += 1
            continue

        
        
        image = _render_pdf_page_to_image(data, index)
        if can_rasterize is None:
            can_rasterize = image is not None
        if image is None:
            note = f"Page {index + 1}: no extractable text and page rendering for OCR is unavailable."
            warnings.append(note)
            results.append(PageResult(index + 1, native_text, "failed", warning=note))
            pages_failed += 1
            continue

        page_warnings: list[str] = []
        ocr_text, method, ocr_status = await _ocr_image_with_fallback(image, page_warnings)
        if is_meaningful(ocr_text):
            results.append(PageResult(index + 1, ocr_text, method, ocr=True))
            pages_ocr += 1
        else:
            if ocr_status == "OCR_UNAVAILABLE":
                ocr_unavailable_pages += 1
                note = f"Page {index + 1}: OCR was unavailable after native text extraction failed."
            else:
                no_readable_pages += 1
                note = f"Page {index + 1} contained no readable text after native extraction and OCR."
            warnings.extend(page_warnings)
            warnings.append(note)
            results.append(PageResult(index + 1, native_text, "failed", warning=note))
            pages_failed += 1

    text = "\n\n".join(f"Page {result.page_number}:\n{result.text}" for result in results if result.text).strip()
    method = "native" if pages_ocr == 0 and pages_failed == 0 else ("mixed" if pages_native else "ocr")
    confidence = round(max(0.0, (pages_native + 0.7 * pages_ocr) / page_count), 2)

    if pages_failed == page_count:
        if can_rasterize is False or ocr_unavailable_pages == pages_failed:
            raise ExtractionError(
                "This PDF could not be read because OCR is unavailable for scanned content.",
                "OCR_UNAVAILABLE",
            )
        if no_readable_pages == pages_failed:
            raise ExtractionError(
                "This PDF was opened successfully, but no readable text was found.",
                "NO_READABLE_CONTENT",
            )
        raise ExtractionError(
            "This PDF could not be read after native extraction and OCR attempts.",
            "UNREADABLE_DOCUMENT",
        )

    metadata = {
        "pages_total": page_count,
        "pages_native": pages_native,
        "pages_ocr": pages_ocr,
        "pages_failed": pages_failed,
    }
    return text, method, warnings, confidence, metadata, results






def _extract_docx(data: bytes) -> tuple[str, list[str]]:
    from docx import Document

    _check_zip_safety(data)
    warnings: list[str] = []
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This document appears to be corrupted.", "CORRUPTED_DOCUMENT") from exc

    parts = []
    try:
        parts.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    except Exception:
        warnings.append("Some paragraphs could not be read.")
    for table_index, table in enumerate(document.tables):
        try:
            parts.extend(" | ".join(str(cell.text).strip() for cell in row.cells) for row in table.rows)
        except Exception:
            warnings.append(f"Table {table_index + 1} could not be fully read.")
    return "\n".join(part for part in parts if part), warnings


def _extract_pptx(data: bytes) -> tuple[str, list[str]]:
    from pptx import Presentation

    _check_zip_safety(data)
    warnings: list[str] = []
    try:
        presentation = Presentation(BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This presentation appears to be corrupted.", "CORRUPTED_DOCUMENT") from exc

    slides = []
    for index, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            try:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows)
            except Exception:
                warnings.append(f"A shape on slide {index} could not be read.")
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        except Exception:
            pass
        slides.append(f"Slide {index}:\n" + "\n".join(parts))
    return "\n\n".join(slides), warnings


def _extract_xlsx(data: bytes) -> tuple[str, list[str]]:
    from openpyxl import load_workbook

    _check_zip_safety(data)
    settings = get_settings()
    warnings: list[str] = []
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError("This spreadsheet appears to be corrupted.", "CORRUPTED_DOCUMENT") from exc

    parts, row_count = [], 0
    for sheet in workbook.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        try:
            for row in sheet.iter_rows(values_only=True):
                if any(value is not None for value in row):
                    row_count += 1
                    if row_count > settings.max_sheet_rows:
                        warnings.append(f"Sheet '{sheet.title}' exceeds the {settings.max_sheet_rows}-row limit; remaining rows were skipped.")
                        break
                    parts.append(" | ".join("" if value is None else str(value) for value in row))
        except Exception:
            warnings.append(f"Sheet '{sheet.title}' could not be fully read.")
            continue
    if not any(line.startswith("Sheet:") for line in parts):
        warnings.append("This workbook has no readable sheets.")
    return "\n".join(parts), warnings


def _extract_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("This text file uses an unsupported encoding.", "UNSUPPORTED_ENCODING")


def _extract_csv(data: bytes) -> tuple[str, list[str]]:
    text = _extract_text(data)
    warnings: list[str] = []
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    try:
        reader = csv.reader(text.splitlines(), dialect)
        for line_number, row in enumerate(reader, 1):
            try:
                rows.append(" | ".join(row))
            except Exception:
                warnings.append(f"Row {line_number} could not be parsed and was skipped.")
    except csv.Error as exc:
        raise ExtractionError("This CSV file could not be parsed.", "CORRUPTED_DOCUMENT") from exc
    return "\n".join(rows), warnings


def _check_zip_safety(data: bytes) -> None:
    settings = get_settings()
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 10_000 or sum(item.file_size for item in infos) > settings.max_zip_uncompressed_bytes:
                raise ExtractionError("This document is too large to process safely.", "DOCUMENT_TOO_LARGE")
    except ExtractionError:
        raise
    except zipfile.BadZipFile as exc:
        raise ExtractionError("This document appears to be corrupted.", "CORRUPTED_DOCUMENT") from exc






async def extract_bytes_async(filename: str | None, data: bytes) -> ExtractionResult:
    name = safe_filename(filename)
    if not data:
        return ExtractionResult(name, "unknown", False, error="The uploaded file is empty.", error_code="EMPTY_FILE")

    try:
        kind = detect_file_type(name, data)
    except ExtractionError as exc:
        return ExtractionResult(name, file_extension(name).lstrip(".") or "unknown", False, error=str(exc), error_code=exc.code)

    try:
        if kind == "image":
            text, method, warnings, confidence, status = await _extract_image_async(data)
            if status == "NO_ANALYZABLE_CONTENT":
                return ExtractionResult(
                    name,
                    kind,
                    True,
                    analyzable=False,
                    status="NO_ANALYZABLE_CONTENT",
                    text="",
                    warnings=warnings,
                    error="Nothing to analyze was found in this image. It appears to be a normal image without readable privacy or document text.",
                    error_code="NO_ANALYZABLE_CONTENT",
                    extraction_method=method,
                    confidence=0.0,
                )
            text, truncated = _limit_text(text)
            return ExtractionResult(
                name,
                kind,
                True,
                analyzable=True,
                status="SUCCESS",
                text=text,
                warnings=warnings,
                extraction_method=method,
                confidence=confidence,
                truncated=truncated,
            )

        if kind == "pdf":
            text, method, warnings, confidence, meta, page_results = await _extract_pdf_async(data)
            text, truncated = _limit_text(text)
            partial = meta.get("pages_failed", 0) > 0
            return ExtractionResult(
                name,
                kind,
                True,
                partial=partial,
                analyzable=True,
                status="SUCCESS",
                text=text,
                pages=page_results,
                metadata=meta,
                warnings=warnings,
                extraction_method=method,
                confidence=confidence,
                truncated=truncated,
            )

        if kind == "docx":
            text, warnings = _extract_docx(data)
        elif kind == "pptx":
            text, warnings = _extract_pptx(data)
        elif kind == "xlsx":
            text, warnings = _extract_xlsx(data)
        elif kind == "csv":
            text, warnings = _extract_csv(data)
        elif kind == "txt":
            text, warnings = _extract_text(data), []
        else:
            return ExtractionResult(name, kind, False, error="This file type is not supported.", error_code="UNSUPPORTED_FILE_TYPE")

        if not text.strip():
            return ExtractionResult(name, kind, False, warnings=warnings, error="No readable content was found in this document.", error_code="EMPTY_DOCUMENT")

        text, truncated = _limit_text(text)
        return ExtractionResult(name, kind, True, partial=bool(warnings), analyzable=True, status="SUCCESS", text=text, warnings=warnings, extraction_method="native", confidence=1.0 if not warnings else 0.85, truncated=truncated)

    except ExtractionError as exc:
        return ExtractionResult(name, kind, False, error=str(exc), error_code=exc.code)
    except Exception as exc:  
        from utils.logger import get_logger

        get_logger(__name__).error("Unexpected extraction failure for %s: %s", name, exc, exc_info=True)
        return ExtractionResult(name, kind, False, error="This file could not be processed.", error_code="EXTRACTION_FAILED")


def extract_file(filename: str | None, data: bytes) -> ExtractionResult:
    import asyncio

    return asyncio.run(extract_bytes_async(filename, data))
